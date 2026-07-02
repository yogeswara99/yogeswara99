import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_border_none(cell):
    pass

def add_bottom_border(paragraph, color="999999", sz=6):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_rich(paragraph, text, size=9.5, color=None, bold_default=False, font="Calibri"):
    """Parse **bold** markers and add runs to an existing paragraph."""
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if not part:
            continue
        bold = bold_default
        if part.startswith('**') and part.endswith('**'):
            part = part[2:-2]
            bold = True
        run = paragraph.add_run(part)
        run.font.size = Pt(size)
        run.font.name = font
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

def new_para(doc, space_before=0, space_after=4, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p

def add_heading(doc, text, navy):
    p = new_para(doc, space_before=10, space_after=5)
    add_rich(p, f"**{text.upper()}**", size=9.5, color=navy)
    add_bottom_border(p, color="CCCCCC", sz=6)
    return p

def add_bullet(doc, text, navy, size=9.7):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Cm(0.45)
    add_rich(p, text, size=size, color="1A1A1A")
    return p

def build_resume(filename, navy, accent, name, title_line, proof_items, summary,
                  competencies, experience_blocks, early_career, skills_line,
                  education, certifications, languages):
    doc = Document()

    # Base style
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(9.7)

    section = doc.sections[0]
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

    # Name
    p = new_para(doc, space_after=2)
    r = p.add_run(name)
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.name = 'Cambria'
    r.font.color.rgb = RGBColor.from_string(navy)

    # Title line
    p = new_para(doc, space_after=4)
    add_rich(p, f"**{title_line}**", size=9.5, color=accent)

    # Contact line
    p = new_para(doc, space_after=6)
    add_rich(p, "Dubai, UAE  |  +971 56 167 8385  |  yogeswarareddy99@gmail.com  |  linkedin.com/in/yogeswara-reddy-gondesi-b2144364",
              size=9, color="555555")
    add_bottom_border(p, color=accent, sz=16)

    # Proof strip
    p = new_para(doc, space_before=6, space_after=8)
    add_rich(p, "   •   ".join(proof_items), size=8.8, color=navy)

    # Summary
    add_heading(doc, "Professional Summary", navy)
    p = new_para(doc, space_after=8)
    add_rich(p, summary, size=9.7)

    # Competencies
    add_heading(doc, "Core Competencies", navy)
    table = doc.add_table(rows=(len(competencies) + 1) // 2, cols=2)
    table.autofit = True
    for i, comp in enumerate(competencies):
        row, col = divmod(i, 2)
        cell = table.cell(row, col)
        cell.paragraphs[0].paragraph_format.space_after = Pt(3)
        add_rich(cell.paragraphs[0], f"-  {comp}", size=9.3)
    new_para(doc, space_after=6)

    # Experience
    add_heading(doc, "Professional Experience", navy)
    for block in experience_blocks:
        p = new_para(doc, space_before=4, space_after=0)
        add_rich(p, f"**{block['company']}**", size=10, color=navy)
        r = p.add_run(f"    {block['location']}")
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = RGBColor.from_string("555555")

        for role in block['roles']:
            p = new_para(doc, space_before=5, space_after=3)
            add_rich(p, f"**{role['title']}**", size=9.7, color=navy)
            r = p.add_run(f"\t{role['dates']}")
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string("555555")
            tab_stops = p.paragraph_format.tab_stops
            tab_stops.add_tab_stop(Cm(16.5))
            for bullet in role['bullets']:
                add_bullet(doc, bullet, navy)

    # Early career
    p = new_para(doc, space_before=8, space_after=4)
    add_rich(p, "**EARLIER CAREER**", size=8, color="777777")
    for line in early_career['lines']:
        p = new_para(doc, space_after=2)
        add_rich(p, line, size=9.3)
    p = new_para(doc, space_after=8)
    add_rich(p, early_career['note'], size=9.3, color="333333")

    # Skills
    add_heading(doc, "Skills & Tools", navy)
    p = new_para(doc, space_after=8)
    add_rich(p, skills_line, size=9.3)

    # Education
    add_heading(doc, "Education", navy)
    for edu in education:
        p = new_para(doc, space_after=0)
        add_rich(p, f"**{edu['degree']}**", size=9.5, color=navy)
        p = new_para(doc, space_after=6)
        add_rich(p, f"{edu['school']}  |  {edu['dates']}", size=9, color="555555")

    # Certifications
    add_heading(doc, "Certifications", navy)
    for cert in certifications:
        add_bullet(doc, cert, navy, size=9.3)

    # Languages
    add_heading(doc, "Languages", navy)
    p = new_para(doc, space_after=4)
    add_rich(p, languages, size=9.3)

    doc.save(filename)
    print("saved", filename)
